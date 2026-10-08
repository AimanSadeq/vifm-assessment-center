"""Build the SDC HiPo IDP development-actions sheet for Ali (internal, VIFM).

Usage: python3 build-idp-sheet.py [out.xlsx]   (default: out/SDC_HiPo_IDP_Actions_v2.xlsx)
One row per SDC competency and target level an IDP can show (25 rows). Goal behaviours come from v3.json;
draft actions, SMART goals, 30/60/90-day milestones and evidence from idp_actions.py. The workbook has no cached formula values; Excel
recalculates on open.
"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.workbook.properties import CalcProperties

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from idp_actions import ACTIONS, MODULE, PLAN  # noqa: E402

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "out", "SDC_HiPo_IDP_Actions_v2.xlsx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)
v3 = json.load(open(os.path.join(HERE, "v3.json")))

NAVY, GREY, YELLOW, LINE = "010131", "F2F2F2", "FFF2CC", "D6DCE5"
def font(**k): return Font(name="Arial", size=k.pop("size", 10), **k)
H_FILL = PatternFill("solid", fgColor=NAVY); Y_FILL = PatternFill("solid", fgColor=YELLOW); G_FILL = PatternFill("solid", fgColor=GREY)
thin = Side(style="thin", color=LINE); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CODE = {"Innovation": "IN", "Results Oriented": "RO", "Effective Collaboration and Communication": "ECC",
        "Professionalism and Integrity": "PI", "Accountability & Ownership": "AO", "Strategic Mindset": "SM",
        "Drives Sustainable Performance": "DSP", "Makes Effective Decisions and Plans": "EDP",
        "Inspires & Motivates": "IM", "Leads by Example": "LBE"}

comps, goals = [], {}
for i in v3["indicators"]:
    key = (i["category"], i["competency"])
    if key not in comps: comps.append(key)
    if i["option_level"] in ("Basic", "Proficient", "Advanced"):
        goals.setdefault((i["competency"], i["option_level"]), []).append(f'{i["id"]}: {i["text"]}')
gh = v3["grade_header"]
def required_at(comp, level):
    col = gh.index(comp)
    gs = sorted(int(g[0]) for g in v3["grades"] if g[col] == level)
    return ", ".join(f"G{g}" for g in gs) if gs else "Not required at any grade; used when building on a strength"

def header(ws, row, labels, widths):
    for c, (lab, w) in enumerate(zip(labels, widths), 1):
        cell = ws.cell(row=row, column=c, value=lab)
        cell.font = font(bold=True, color="FFFFFF"); cell.fill = H_FILL; cell.border = BOX
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = w

wb = Workbook()

# Read me
rm = wb.active; rm.title = "Read me"
rm.column_dimensions["A"].width = 24; rm.column_dimensions["B"].width = 110
lines = [
    ("title", "SDC HiPo assessment: IDP development actions"),
    ("sub", "Version 2, 8 October 2026 (adds SMART goals, 30/60/90-day milestones and specific evidence, per Ali AlShouli's review). INTERNAL to VIFM. For Ali to review and complete by Wednesday 14 October 2026."),
    ("h", "Purpose"),
    ("p", "What this is", "The development actions Caliber prints on each participant's one-page IDP. One row per SDC competency and target level an IDP can show (25 rows)."),
    ("p", "How Caliber uses it", "For a competency below the grade requirement, the IDP shows the row for the REQUIRED level. For the strength chosen to build on, it shows the row for the NEXT level up. Each IDP shows two gaps to close and one strength to build."),
    ("p", "Each priority shows", "The goal behaviours (from v3), one SMART goal, milestones at 30, 60 and 90 days, the actions (on the job, coaching, learning) and the evidence that shows progress."),
    ("p", "Drafts", "Every yellow cell is pre-filled with a VIFM draft so you can edit rather than start from blank. Rewrite anything that does not fit SDC, then set Status to 'Approved'."),
    ("h", "What to fill in"),
    ("p", "Yellow cells", "Programme module, SMART goal, the three milestones, On the job, Coaching, Learning, Evidence of progress, Status and Notes. Grey cells are fixed (from the v3 workbook)."),
    ("p", "70-20-10 approach", "A guide to the mix of development activities, not a scoring formula. On the job = an assignment or practice in the participant's real work. Coaching = what to work on with the VIFM coach. Learning = the programme module session or activity that builds it."),
    ("h", "Writing rules"),
    ("p", "1. Specific", "One concrete action per cell, starting with a verb, about 15 to 30 words."),
    ("p", "SMART goal", "Specific, measurable, achievable, relevant and time-bound: one outcome the participant can show by day 90, with how it will be measured."),
    ("p", "Milestones", "What is done by day 30, 60 and 90, so the coach and SDC can see whether the participant is on track."),
    ("p", "Evidence", "Specific proof, such as a completed project or output, a documented decision, or line manager or stakeholder feedback. Avoid general 'observed behaviours'."),
    ("p", "2. Doable", "Achievable within the six-month programme and the participant's current role."),
    ("p", "3. Level-appropriate", "Basic actions stay within the participant's own work; Proficient ones involve the team and stakeholders; Advanced ones have organisational or strategic reach (matching the v3 level definitions)."),
    ("p", "4. Linked", "Tie the learning action to the module that develops the competency, so the IDP connects to the programme."),
    ("p", "5. Line managers", "Line manager feedback is used as evidence; confirm with SDC how line managers will take part."),
    ("h", "Timeline"),
    ("p", "Wed 14 Oct", "All 25 rows reviewed and set to 'Approved'."),
    ("p", "Thu 15 Oct", "Loaded into Caliber; you review the first real IDPs."),
    ("h", "Reference"),
    ("p", "70-20-10", "Lombardo, M. M., & Eichinger, R. W. (1996). The Career Architect Development Planner. Minneapolis: Lominger. Center for Creative Leadership, The 70-20-10 Rule for Leadership Development (ccl.org). The programme itself is built on 10/20/70 (proposal 20260809-SDC-02-HPHLD)."),
    ("p", "SMART goals", "Doran, G. T. (1981). There's a S.M.A.R.T. way to write management's goals and objectives. Management Review, 70(11), 35-36."),
]
r = 1
for item in lines:
    kind = item[0]
    if kind == "title": rm.cell(row=r, column=1, value=item[1]).font = font(size=16, bold=True, color=NAVY)
    elif kind == "sub": rm.cell(row=r, column=1, value=item[1]).font = font(italic=True, color="7F7F7F")
    elif kind == "h":
        r += 1
        c = rm.cell(row=r, column=1, value=item[1]); c.font = font(size=12, bold=True, color="FFFFFF"); c.fill = H_FILL
        rm.cell(row=r, column=2).fill = H_FILL
    else:
        a = rm.cell(row=r, column=1, value=item[1]); a.font = font(bold=True); a.alignment = WRAP
        b = rm.cell(row=r, column=2, value=item[2]); b.font = font(); b.alignment = WRAP
        if item[1] == "Yellow cells": a.fill = Y_FILL
    r += 1
rm.sheet_view.showGridLines = False

# Actions
ws = wb.create_sheet("IDP actions")
ws["A1"] = "SDC HiPo: IDP development actions (drafts for Ali to review)"; ws["A1"].font = font(size=14, bold=True, color=NAVY)
ws["A2"] = "Yellow cells are editable drafts; grey cells come from the v3 workbook. Set Status to 'Approved' once a row is final."
ws["A2"].font = font(italic=True, color="7F7F7F")
cols = ["Ref", "Category", "SDC competency", "Target level", "Required at grades", "Goal behaviours (v3 indicators)",
        "Programme module", "SMART goal (by day 90)", "By day 30", "By day 60", "By day 90",
        "On the job", "Coaching", "Learning", "Evidence of progress", "Status", "Notes"]
header(ws, 4, cols, [11, 11, 24, 11, 16, 52, 20, 40, 30, 30, 30, 40, 34, 34, 36, 14, 26])
rr = 5
for cat, comp in comps:
    code = CODE[comp]
    for lvl in ("Basic", "Proficient", "Advanced"):
        # Only rows an IDP can show: a level required at some grade (a gap target), or the
        # level above a requirement (a strength to build). Core is never required at Basic.
        if lvl == "Basic" and required_at(comp, lvl).startswith("Not required"):
            continue
        job, coach, learn = ACTIONS[code][lvl]
        smart, d30, d60, d90, evidence = PLAN[code][lvl]
        fixed = [f"IDP-{code}-{lvl[0]}", cat, comp, lvl, required_at(comp, lvl), "\n".join(goals[(comp, lvl)])]
        editable = [MODULE[code], smart, d30, d60, d90, job, coach, learn, evidence, "Draft", None]
        for c, v in enumerate(fixed + editable, 1):
            cell = ws.cell(row=rr, column=c, value=v); cell.font = font(); cell.alignment = WRAP; cell.border = BOX
            cell.fill = G_FILL if c <= 6 else Y_FILL
        ws.row_dimensions[rr].height = 120
        rr += 1
last = rr - 1
dv = DataValidation(type="list", formula1='"Draft,Ali reviewed,Approved"', allow_blank=False); ws.add_data_validation(dv); dv.add(f"P5:P{last}")
dvm = DataValidation(type="list", formula1='"M1 Leading Self,M2 Leading Others,M3 Leading the Business,M4 Leading for the Future"', allow_blank=False)
ws.add_data_validation(dvm); dvm.add(f"G5:G{last}")
ws.conditional_formatting.add(f"P5:P{last}", CellIsRule(operator="equal", formula=['"Approved"'], fill=PatternFill("solid", fgColor="E2EFDA")))
ws["N1"] = "Approved"; ws["N1"].font = font(bold=True)
ws["O1"] = f'=COUNTIF(P5:P{last},"Approved")&" of "&COUNTA(A5:A{last})'; ws["O1"].font = font()
ws.freeze_panes = "E5"; ws.auto_filter.ref = f"A4:Q{last}"

# Programme modules
pm = wb.create_sheet("Programme modules")
header(pm, 1, ["Module", "Focus (from the programme plan)", "Main competencies linked"], [26, 80, 50])
mods = [
    ("M1 Leading Self", "Self-awareness, strengths, emotional intelligence, personal leadership, goal setting; output: Individual Development Plan."),
    ("M2 Leading Others", "Team leadership, coaching, feedback, communication, influence, collaboration; output: team leadership action plan."),
    ("M3 Leading the Business", "Strategy, business and financial acumen, decision-making, innovation, problem solving; output: business challenge and simulation."),
    ("M4 Leading for the Future", "Change and stakeholder management, executive presence, succession readiness; output: capstone presentation and readiness plan."),
]
for k, (m, f) in enumerate(mods, 2):
    linked = ", ".join(comp for _, comp in comps if MODULE[CODE[comp]] == m)
    for c, v in enumerate([m, f, linked], 1):
        cell = pm.cell(row=k, column=c, value=v); cell.font = font(); cell.alignment = WRAP; cell.border = BOX
pm.cell(row=7, column=1, value="Source: SDC HiPo project charter V0.1 (15 Sep 2026) and proposal 20260809-SDC-02-HPHLD, Part 2.").font = font(italic=True, color="7F7F7F")

wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
print("saved", OUT)
