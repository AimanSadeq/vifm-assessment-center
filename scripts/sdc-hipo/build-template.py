"""Build the SDC HiPo scenario-question template for Ali (internal, VIFM).

Usage: python3 build-template.py [out.xlsx]   (default: out/SDC_HiPo_Scenario_Template_v1.xlsx)
Reads v3.json (Ali's v3 indicators + SDC grade mapping) and bank.json (closest items from
docs/competency-items-*.md). The workbook has no cached formula values; Excel recalculates on open.
"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.comments import Comment

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "out", "SDC_HiPo_Scenario_Template_v1.xlsx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)
v3 = json.load(open(os.path.join(HERE, "v3.json")))
bank = json.load(open(os.path.join(HERE, "bank.json")))
IND = v3["indicators"]
N_IND = len(IND)

NAVY = "010131"; GREY = "F2F2F2"; YELLOW = "FFF2CC"; LINE = "D6DCE5"
F = "Arial"
def font(**k): return Font(name=F, size=k.pop("size", 10), **k)
H_FILL = PatternFill("solid", fgColor=NAVY)
Y_FILL = PatternFill("solid", fgColor=YELLOW)
G_FILL = PatternFill("solid", fgColor=GREY)
thin = Side(style="thin", color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
BLUE = "0000FF"

def header(ws, row, labels, widths=None):
    for c, lab in enumerate(labels, 1):
        cell = ws.cell(row=row, column=c, value=lab)
        cell.font = font(bold=True, color="FFFFFF"); cell.fill = H_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center"); cell.border = BOX
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=row, column=c).column_letter].width = w

PREFIX = {"Innovation": "IN", "Results Oriented": "RO", "Effective Collaboration and Communication": "ECC",
          "Professionalism and Integrity": "PI", "Accountability & Ownership": "AO", "Strategic Mindset": "SM",
          "Drives Sustainable Performance": "DSP", "Makes Effective Decisions and Plans": "EDP",
          "Inspires & Motivates": "IM", "Leads by Example": "LBE"}
comps = []
for i in IND:
    if (i["category"], i["competency"]) not in comps: comps.append((i["category"], i["competency"]))
gh = v3["grade_header"]; grades = sorted(v3["grades"], key=lambda r: int(r[0]))
def req_text(comp):
    col = gh.index(comp)
    return "; ".join(f"G{g[0]} {g[col]}" for g in grades)

wb = Workbook()

# ---------------- Read me ----------------
rm = wb.active; rm.title = "Read me"
rm.column_dimensions["A"].width = 26; rm.column_dimensions["B"].width = 110
rows = [
    ("title", "SDC HiPo assessment: scenario question template", None),
    ("sub", "Version 1, 8 October 2026. INTERNAL to VIFM. Do not share with candidates or paste into public AI tools.", None),
    ("h", "Purpose", None),
    ("p", "What this is", "The workbook for writing the 30 scenario questions (3 per SDC competency) that form Part 1 of the SDC HiPo assessment. Part 2 is the timed Logica reasoning tests. Total time limit agreed with SDC: 2 hours."),
    ("p", "Content basis", "SDC competency names and grade requirements (SDC mapping, Grade 4 = Grade 5 per Suzanne, 6 Oct 2026) and the 160 indicators in SDC_HiPo_Caliber_Detailed_Assessor_Indicators_v3.xlsx (sheets 'Indicators (v3)' and 'Grade requirements' here)."),
    ("h", "What you fill in", None),
    ("p", "Yellow cells", "Sheet 'Scenarios': the situation, the four responses, what level each response reflects, its v3 indicator ID, and the status. Grey cells are fixed or calculated: do not edit them."),
    ("p", "Blue text", "Assumptions you can change (sheets 'Time budget' and 'Scoring (proposed)')."),
    ("p", "Worked examples", "Sheet 'Worked examples' has three complete Innovation scenarios (Basic, Proficient, Advanced) in this exact format. You may adopt them as the Innovation rows or replace them."),
    ("p", "Starting points", "Sheet 'Existing VIFM items' gives the closest item from Caliber's existing scenario bank for each SDC competency. They are generic and not level-keyed: adapt the situation to SDC and rewrite the responses to one per level."),
    ("h", "How each scenario works", None),
    ("p", "Candidate sees", "A short workplace situation and four responses, then answers: 'Which response is the MOST effective? Which is the LEAST effective?' Response order is shuffled for each candidate."),
    ("p", "Four responses", "Each response reflects exactly one level, and each level appears once per scenario: Advanced, Proficient, Basic, and Counter-evidence (a common mistake, from the negative indicators). Link each response to one v3 indicator ID of the same competency and the same level."),
    ("p", "Scenario level", "Sets the scope of the situation (from the v3 definitions). Basic: own work or immediate area, straightforward. Proficient: varied situations, team and stakeholders. Advanced: complex situations with organizational or strategic impact."),
    ("p", "Checks", "Columns 'Level check' and 'Indicator check' on the Scenarios sheet must both read OK before you set the status to 'Ready for review'."),
    ("h", "Writing rules", None),
    ("p", "1. Context", "Set situations in SDC's world (destination development, tourism and visitor experience, heritage, environment, and corporate functions such as finance, procurement, HR, IT). Never name real people, projects or confidential data."),
    ("p", "2. Length", "Situation 50 to 100 words (see 'Situation words'), with one clear dilemma and enough detail to judge."),
    ("p", "3. Balance", "Four responses of similar length (within about 10 words) and tone. All must be plausible; no obviously silly option."),
    ("p", "4. Best response", "The Advanced response must genuinely be the most effective in that situation, not an over-engineered answer."),
    ("p", "5. Counter-evidence", "The Counter-evidence response should be tempting: a mistake people really make, matching one negative indicator."),
    ("p", "6. No give-aways", "Do not let wording reveal the level (for example, putting 'strategic' or 'stakeholders' only in the best response)."),
    ("p", "7. One behaviour", "One clear action per response. Avoid 'always' and 'never'. Plain English, no jargon."),
    ("p", "8. Fairness", "Gender-neutral, culturally appropriate for KSA, no knowledge of SDC internals needed to answer."),
    ("p", "9. Stand-alone", "Each scenario stands alone; do not refer to other scenarios or to 'option A', as order is shuffled."),
    ("h", "Scoring (proposed)", None),
    ("p", "Method", "See sheet 'Scoring (proposed)': points for the level of the response chosen as MOST effective, plus a point for choosing the Counter-evidence response as LEAST effective. Each competency gets a level from its 3 scenarios, compared with the required level for the candidate's grade. Final scoring rules are set after internal review and the pilot."),
    ("p", "Reporting", "Results are reported as judgement in realistic situations, not as observed behaviour."),
    ("h", "Timeline", None),
    ("p", "Sun 11 Oct", "One sample scenario per competency (10 in total), status 'Ready for review', to Aiman. After review they go to Suzanne as the samples she asked for."),
    ("p", "Wed 14 Oct", "All 30 scenarios complete; SDC sign-off on scenarios and sample reports."),
    ("p", "Sun 18 Oct", "Pilot with 3 to 5 internal testers."),
    ("p", "Tue 20 Oct", "Launch, only after SDC's green light."),
    ("h", "References", None),
    ("p", "Response instructions", "McDaniel, M. A., Hartman, N. S., Whetzel, D. L., & Grubb, W. L. (2007). Situational judgment tests, response instructions, and validity: A meta-analysis. Personnel Psychology, 60(1), 63-91."),
    ("p", "SJT design", "Whetzel, D. L., & McDaniel, M. A. (2009). Situational judgment tests: An overview of current research. Human Resource Management Review, 19(3), 188-202."),
    ("p", "Caliber bank", "Existing VIFM scenario items: docs/competency-items-cluster1.md and docs/competency-items-clusters2-8.md in the Caliber repository (MOST/LEAST format, expert-keyed)."),
]
r = 1
for kind, a, b in rows:
    if kind == "title":
        c = rm.cell(row=r, column=1, value=a); c.font = font(size=16, bold=True, color=NAVY)
    elif kind == "sub":
        c = rm.cell(row=r, column=1, value=a); c.font = font(italic=True, color="7F7F7F")
    elif kind == "h":
        r += 1
        c = rm.cell(row=r, column=1, value=a); c.font = font(size=12, bold=True, color="FFFFFF"); c.fill = H_FILL
        rm.cell(row=r, column=2).fill = H_FILL
    else:
        c1 = rm.cell(row=r, column=1, value=a); c1.font = font(bold=True); c1.alignment = WRAP
        c2 = rm.cell(row=r, column=2, value=b); c2.font = font(); c2.alignment = WRAP
        if a == "Yellow cells": c1.fill = Y_FILL
        if a == "Blue text": c1.font = font(bold=True, color=BLUE)
    r += 1

# ---------------- Indicators (v3) ----------------
iv = wb.create_sheet("Indicators (v3)")
header(iv, 1, ["Category", "SDC competency", "v3 proficiency level", "Response level", "Indicator ID", "Observable behavioural indicator", "Evidence type"],
       [12, 34, 24, 16, 12, 80, 26])
for k, i in enumerate(IND, 2):
    for c, v in enumerate([i["category"], i["competency"], i["level"], i["option_level"], i["id"], i["text"], i["evidence"]], 1):
        cell = iv.cell(row=k, column=c, value=v); cell.font = font(); cell.alignment = WRAP; cell.border = BOX
iv.freeze_panes = "A2"; iv.auto_filter.ref = f"A1:G{N_IND+1}"
LAST = N_IND + 1
IND_ID = f"'Indicators (v3)'!$E$2:$E${LAST}"
IND_COMP = f"'Indicators (v3)'!$B$2:$B${LAST}"
IND_LVL = f"'Indicators (v3)'!$D$2:$D${LAST}"
IND_TXT = f"'Indicators (v3)'!$F$2:$F${LAST}"

# ---------------- Grade requirements ----------------
gr = wb.create_sheet("Grade requirements")
header(gr, 1, gh, [8, 12, 26, 34] + [16] * (len(gh) - 4))
for k, g in enumerate(sorted(v3["grades"], key=lambda r: -int(r[0])), 2):
    for c, v in enumerate(g, 1):
        cell = gr.cell(row=k, column=c, value=int(v) if c == 1 else v); cell.font = font(); cell.alignment = WRAP; cell.border = BOX
n = len(v3["grades"]) + 3
gr.cell(row=n, column=1, value="Source: SDC positional mapping (grades 5 to 8); Grade 4 uses Grade 5 levels, confirmed by Suzanne Salim by email on 6 Oct 2026. Leadership competencies are required only at Grades 7 and 8.").font = font(italic=True, color="7F7F7F")

# ---------------- Scoring (proposed) ----------------
sc = wb.create_sheet("Scoring (proposed)")
sc.column_dimensions["A"].width = 30; sc.column_dimensions["B"].width = 14; sc.column_dimensions["C"].width = 90
header(sc, 1, ["Response chosen as MOST effective", "Points", "Note"])
pts = [("Advanced", 3), ("Proficient", 2), ("Basic", 1), ("Counter-evidence", 0)]
for k, (lvl, p) in enumerate(pts, 2):
    sc.cell(row=k, column=1, value=lvl).font = font()
    c = sc.cell(row=k, column=2, value=p); c.font = font(color=BLUE)
    sc.cell(row=k, column=3, value="Proposed weight; confirm after internal review and pilot.").font = font(color="7F7F7F")
header(sc, 7, ["Response chosen as LEAST effective", "Points", "Note"])
sc.cell(row=8, column=1, value="Counter-evidence").font = font()
c = sc.cell(row=8, column=2, value=1); c.font = font(color=BLUE)
sc.cell(row=8, column=3, value="Recognising the common mistake earns a point; any other LEAST choice earns 0.").font = font(color="7F7F7F")
header(sc, 10, ["Per scenario", "Points", "Note"])
sc.cell(row=11, column=1, value="Maximum points per scenario").font = font()
sc.cell(row=11, column=2, value="=MAX(B2:B5)+B8").font = font()
sc.cell(row=11, column=3, value="Calculated.").font = font(color="7F7F7F")
sc.cell(row=12, column=1, value="Maximum per competency (3 scenarios)").font = font()
sc.cell(row=12, column=2, value="=3*B11").font = font()
sc.cell(row=12, column=3, value="Calculated. Cut-offs from points to Below Basic / Basic / Proficient / Advanced are set during the build and checked in the pilot.").font = font(color="7F7F7F")
for rr in range(2, 13):
    for cc in range(1, 4): sc.cell(row=rr, column=cc).alignment = WRAP
PTS_LVL = "'Scoring (proposed)'!$A$2:$A$5"; PTS_VAL = "'Scoring (proposed)'!$B$2:$B$5"

# ---------------- Scenarios ----------------
SC_COLS = ["Scenario ID", "Category", "SDC competency", "Scenario level", "Required level by grade",
           "Situation (50 to 100 words)",
           "Response A", "A reflects", "A indicator ID",
           "Response B", "B reflects", "B indicator ID",
           "Response C", "C reflects", "C indicator ID",
           "Response D", "D reflects", "D indicator ID",
           "Situation words", "Level check", "Indicator check", "Status", "Reviewer notes"]
SC_W = [13, 11, 24, 11, 26, 48, 34, 15, 12, 34, 15, 12, 34, 15, 12, 34, 15, 12, 10, 16, 18, 15, 30]
LEVELS = '"Advanced,Proficient,Basic,Counter-evidence"'

def scenario_sheet(ws, first_row, rows):
    header(ws, first_row - 1, SC_COLS, SC_W)
    ws.row_dimensions[first_row - 1].height = 30
    for k, row in enumerate(rows):
        rr = first_row + k
        for c in range(1, len(SC_COLS) + 1):
            cell = ws.cell(row=rr, column=c); cell.font = font(); cell.alignment = WRAP; cell.border = BOX
        for c, v in enumerate(row["fixed"], 1):
            ws.cell(row=rr, column=c, value=v).fill = G_FILL
        for c in range(6, 19):
            ws.cell(row=rr, column=c).fill = Y_FILL
        ws.cell(row=rr, column=22).fill = Y_FILL
        for c, v in row.get("fill", {}).items():
            ws.cell(row=rr, column=c, value=v)
        ws.cell(row=rr, column=19, value=f'=IF(F{rr}="","",LEN(TRIM(F{rr}))-LEN(SUBSTITUTE(TRIM(F{rr})," ",""))+1)')
        lv = [f"H{rr}", f"K{rr}", f"N{rr}", f"Q{rr}"]
        each = ",".join(f'({"+".join(f"({x}={chr(34)}{L}{chr(34)})" for x in lv)})=1' for L in ["Advanced", "Proficient", "Basic", "Counter-evidence"])
        ws.cell(row=rr, column=20, value=f'=IF(COUNTA({",".join(lv)})<4,"Pending",IF(AND({each}),"OK","Use each level once"))')
        pairs = [("H", "I"), ("K", "L"), ("N", "O"), ("Q", "R")]
        oks = ",".join(f'INDEX({IND_COMP},MATCH({i}{rr},{IND_ID},0))=$C{rr},INDEX({IND_LVL},MATCH({i}{rr},{IND_ID},0))={l}{rr}' for l, i in pairs)
        ids = ",".join(f"{i}{rr}" for _, i in pairs)
        ws.cell(row=rr, column=21, value=f'=IF(COUNTA({ids})<4,"Pending",IFERROR(IF(AND({oks}),"OK","Check competency or level"),"Unknown ID"))')
        for c in (19, 20, 21): ws.cell(row=rr, column=c).fill = G_FILL
        ws.row_dimensions[rr].height = 120
    last = first_row + len(rows) - 1
    dv_l = DataValidation(type="list", formula1=LEVELS, allow_blank=True)
    dv_i = DataValidation(type="list", formula1=IND_ID, allow_blank=True)
    dv_s = DataValidation(type="list", formula1='"Draft,Ready for review,Revise,Approved"', allow_blank=True)
    for d in (dv_l, dv_i, dv_s): ws.add_data_validation(d)
    for col in "HKNQ": dv_l.add(f"{col}{first_row}:{col}{last}")
    for col in "ILOR": dv_i.add(f"{col}{first_row}:{col}{last}")
    dv_s.add(f"V{first_row}:V{last}")
    green = PatternFill("solid", fgColor="E2EFDA"); amber = PatternFill("solid", fgColor="FCE4D6")
    for col in "TU":
        rng = f"{col}{first_row}:{col}{last}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OK"'], fill=green))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND({col}{first_row}<>"OK",{col}{first_row}<>"Pending")'], fill=amber))
    ws.freeze_panes = ws.cell(row=first_row, column=4)
    return last

s = wb.create_sheet("Scenarios", 1)
s["A1"] = "SDC HiPo: scenario questions (Part 1, untimed)"; s["A1"].font = font(size=14, bold=True, color=NAVY)
s["A2"] = "Candidates answer for each scenario: 'Which response is the MOST effective? Which is the LEAST effective?' Fill the yellow cells; grey cells are fixed or calculated."
s["A2"].font = font(italic=True, color="7F7F7F")
rows = []
for cat, comp in comps:
    for lvl in ("Basic", "Proficient", "Advanced"):
        rows.append({"fixed": [f"SDC-{PREFIX[comp]}-{lvl[0]}", cat, comp, lvl, req_text(comp)]})
last = scenario_sheet(s, 5, rows)
s["F2"] = None
s["P1"] = "Ready for review"; s["P1"].font = font(bold=True)
s["Q1"] = f'=COUNTIF(V5:V{last},"Ready for review")'
s["P2"] = "Approved"; s["P2"].font = font(bold=True)
s["Q2"] = f'=COUNTIF(V5:V{last},"Approved")'
s["R1"] = "of"; s["S1"] = f"=COUNTA(A5:A{last})"
for ref in ("Q1", "Q2", "R1", "S1"): s[ref].font = font()
SC_LAST = last

# ---------------- Worked examples ----------------
ex = wb.create_sheet("Worked examples", 2)
ex["A1"] = "Worked examples: Innovation (draft, written for SDC)"; ex["A1"].font = font(size=14, bold=True, color=NAVY)
ex["A2"] = "Complete scenarios in the exact template format. Ali may adopt them as the Innovation rows on 'Scenarios' or replace them. Below the examples, each response is explained with its indicator and proposed score."
ex["A2"].font = font(italic=True, color="7F7F7F")
EX = [
    ("Basic",
     "You prepare a weekly visitor-feedback summary for your department. Copying comments from three separate systems into one spreadsheet takes most of Thursday, and colleagues often ask you for the same figures in different formats. Nobody has asked you to change the process, and the summary is always delivered on time.",
     [("Build a simple template that combines the three exports, use it for a month and track how much time it saves.", "Proficient", "IN-P3"),
      ("Mention to your manager that the process is slow and suggest that it could probably be simplified.", "Basic", "IN-B1"),
      ("Pilot a small shared dashboard with the colleagues who use the summary, check time saved and accuracy, then offer it to other departments.", "Advanced", "IN-A3"),
      ("Keep the current method, since the summary is always delivered on time and nobody has complained.", "Counter-evidence", "IN-N1")]),
    ("Proficient",
     "Your team manages bookings for guided heritage tours. Cancellations have risen for three months in a row. A colleague proposes adopting the online booking flow used by another destination. Your manager is sceptical and says the current process has always worked, but has asked for your view before next week's team meeting.",
     [("Agree with your manager that the current process works and that the rise is probably seasonal.", "Counter-evidence", "IN-N2"),
      ("Compare the proposed flow and two other options against the cancellation data, then recommend one with the evidence.", "Proficient", "IN-P2"),
      ("Propose a time-limited trial linked to visitor-experience targets, with success measures agreed with operations and a plan to extend it if it works.", "Advanced", "IN-A2"),
      ("Tell your manager you are open to trying the new flow if it is likely to reduce cancellations.", "Basic", "IN-B3")]),
    ("Advanced",
     "Several departments are piloting their own digital tools for visitor services. Costs overlap, the tools do not share data, and some pilots have run for over a year without a decision. The executive team asks you to recommend how the organization should manage innovation in visitor services over the next two years.",
     [("Select the most advanced tool on the market and require all departments to adopt it from next quarter.", "Counter-evidence", "IN-N3"),
      ("Collect lessons learned from each department's pilot and circulate them in a summary note.", "Basic", "IN-B4"),
      ("Review the current pilots, keep the two with the clearest results and ask for monthly progress reports.", "Proficient", "IN-P3"),
      ("Set up a shared process: a small fund for time-limited pilots, common success measures tied to priorities, and a board that scales or stops each pilot.", "Advanced", "IN-A1")]),
]
ex_rows = []
for lvl, sit, opts in EX:
    fill = {6: sit}
    for k, (txt, lv, iid) in enumerate(opts):
        base = 7 + 3 * k
        fill[base] = txt; fill[base + 1] = lv; fill[base + 2] = iid
    fill[22] = "Ready for review"
    ex_rows.append({"fixed": [f"SDC-IN-{lvl[0]}", "Core", "Innovation", lvl, req_text("Innovation")], "fill": fill})
ex_last = scenario_sheet(ex, 5, ex_rows)
for rr in range(5, ex_last + 1):
    for c in range(6, 19): ex.cell(row=rr, column=c).fill = PatternFill(fill_type=None)

# Explanation table
t0 = ex_last + 3
ex.cell(row=t0 - 1, column=1, value="How each example response is keyed and scored (proposed)").font = font(size=12, bold=True, color=NAVY)
header(ex, t0, ["Scenario ID", "Response", "Reflects", "Indicator ID", "Indicator (from v3)", "Points if chosen as MOST", "Points if chosen as LEAST"])
rr = t0 + 1
for k, (lvl, sit, opts) in enumerate(EX):
    srow = 5 + k
    for j, letter in enumerate("ABCD"):
        lcol = ["H", "K", "N", "Q"][j]; icol = ["I", "L", "O", "R"][j]
        vals = [f"=A{srow}", letter, f"={lcol}{srow}", f"={icol}{srow}",
                f"=IFERROR(INDEX({IND_TXT},MATCH(D{rr},{IND_ID},0)),\"\")",
                f"=IFERROR(INDEX({PTS_VAL},MATCH(C{rr},{PTS_LVL},0)),\"\")",
                f"=IF(C{rr}=\"Counter-evidence\",'Scoring (proposed)'!$B$8,0)"]
        for c, v in enumerate(vals, 1):
            cell = ex.cell(row=rr, column=c, value=v); cell.font = font(); cell.alignment = WRAP; cell.border = BOX; cell.fill = G_FILL
        rr += 1
ex.cell(row=rr + 1, column=1, value="Best answer pattern: MOST = the Advanced response, LEAST = the Counter-evidence response, for 4 points per scenario under the proposed weights.").font = font(italic=True, color="7F7F7F")

# ---------------- Existing VIFM items ----------------
eb = wb.create_sheet("Existing VIFM items", 3)
eb["A1"] = "Closest items in Caliber's existing scenario bank (starting points only)"; eb["A1"].font = font(size=14, bold=True, color=NAVY)
eb["A2"] = "Generic workplace items in MOST/LEAST format, keyed to one best and one worst response, not to levels. To use one: set it in SDC's context and rewrite the four responses as Advanced, Proficient, Basic and Counter-evidence."
eb["A2"].font = font(italic=True, color="7F7F7F")
header(eb, 4, ["SDC competency", "Closest VIFM competency", "Situation", "A", "B", "C", "D", "Key"], [26, 24, 46, 34, 34, 34, 34, 18])
for k, (comp, it) in enumerate(bank.items(), 5):
    st = it["stem"].replace(" MOST / LEAST:", "").replace("MOST / LEAST:", "").strip()
    vals = [comp, it["name"], st] + [it["options"].get(x, "") for x in "ABCD"] + [it["key"]]
    for c, v in enumerate(vals, 1):
        cell = eb.cell(row=k, column=c, value=v); cell.font = font(); cell.alignment = WRAP; cell.border = BOX
    eb.row_dimensions[k].height = 75
eb.freeze_panes = "C5"

# ---------------- Time budget ----------------
tb = wb.create_sheet("Time budget")
tb.column_dimensions["A"].width = 44; tb.column_dimensions["B"].width = 12; tb.column_dimensions["C"].width = 90
header(tb, 1, ["Item", "Minutes", "Basis"])
items = [
    ("Number of scenarios", f"=COUNTA(Scenarios!A5:A{SC_LAST})", "Calculated from the Scenarios sheet (3 per competency).", False),
    ("Minutes per scenario", 2, "Assumption for reading a short situation and choosing MOST and LEAST; confirm in the pilot.", True),
    ("Part 1: scenario questions", "=B2*B3", "Calculated. Untimed for candidates; this is the expected time.", False),
    ("Part 2: Logica reasoning tests (timed)", 40, "Caliber's default Logica time limit (src/lib/assessment-timers.ts, 9 items per subtest). Confirm which subtests SDC uses.", True),
    ("Welcome, consent and demographic questions", 10, "Assumption; SDC is providing the welcome message and the demographic fields.", True),
    ("Total expected time", "=B4+B5+B6", "Calculated.", False),
    ("Limit agreed with SDC", 120, "SDC requirement, from Ahmad Rashid's meeting with Suzanne Salim on 7 Oct 2026.", True),
    ("Headroom", "=B8-B7", "Calculated. Must not be negative.", False),
    ("Status", '=IF(B7<=B8,"Within limit","Over limit")', "Calculated.", False),
]
for k, (a, b, c, is_input) in enumerate(items, 2):
    tb.cell(row=k, column=1, value=a).font = font(bold=a in ("Total expected time", "Status"))
    cb = tb.cell(row=k, column=2, value=b); cb.font = font(color=BLUE if is_input else "000000")
    tb.cell(row=k, column=3, value=c).font = font(color="7F7F7F")
    for cc in range(1, 4): tb.cell(row=k, column=cc).border = BOX; tb.cell(row=k, column=cc).alignment = WRAP

wb.move_sheet("Indicators (v3)", offset=2)
for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False if ws.title == "Read me" else True
from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
print("saved", OUT)
