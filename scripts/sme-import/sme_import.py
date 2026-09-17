"""
SME review importer - loads returned SME validation workbooks back into the live banks.

Every SME workbook carries each item's database id in a grey "Item ID (do not edit)"
column. This script reads a returned workbook, matches every reviewed row to the live
bank by that id, and works out exactly what should change. It never guesses:

  * PREVIEW (default) writes nothing. It prints a summary and saves an Excel report
    listing every change it would make and every row it is holding back, with why.
  * --apply performs the changes, then records every decision (applied or held) in the
    immutable sme_review_log table (migration 00203).

What gets applied automatically
  Approve   -> the item's review status becomes approved, with the reviewer's name and date.
  Reject    -> the item's review status becomes rejected.
  Revise    -> the replacement wording is written, when the reviewer supplied it in full
               (and in Arabic too, where the item has Arabic).
  New AC rating anchors typed into an Assessment Center workbook are inserted.
  Technical Cut-score sheet -> the domain's pass mark (tech_assessment_cut_scores).

What is always HELD for a human (logged, never written)
  * the item id is missing from the bank, or appears twice in the workbook
  * the bank text changed after the workbook was issued (the reviewer saw an old version)
  * a verdict contradicts the reviewer's own checks (Approve with "Key correct? = No")
  * a free-text fix that needs a person to edit options or keys (Technical, Logica, scenarios)
  * English revised without matching Arabic, on an item that serves Arabic
  * a rejection that would break a live test (a Logica facet/difficulty cell with no items
    left, a Persona competency with fewer than 3 statements or no reverse-keyed statement)
  * AR Compass wording changes: the live bank (v1.1) is pinned by running assessments, so
    every rewording is queued for AR Compass v1.2 (decision 2026-09-17). Its held_reason starts
    with ARC_V12_QUEUE, which is how the v1.2 build finds them in sme_review_log.

Usage
  python scripts/sme-import/sme_import.py RETURNED.xlsx [MORE.xlsx | FOLDER ...]
         [--reviewer "Name"] [--apply] [--report PATH] [--reimport]

The reviewer defaults to the "Assigned to" name for the file in the master register
(.tmp/sme/VIFM-SME-assignment-register.xlsx) when present; pass --reviewer otherwise.
Needs NEXT_PUBLIC_SUPABASE_URL and SUPABASE_SERVICE_ROLE_KEY in .env.local.
"""
from __future__ import annotations

import argparse
import getpass
import hashlib
import os
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
REGISTER = ROOT / ".tmp" / "sme" / "VIFM-SME-assignment-register.xlsx"
ID = "Item ID (do not edit)"
VERDICT = "SME verdict"
TIP = "[DEV TIP]"
ARC_V12_QUEUE = "Queued for AR Compass v1.2"
NOW = datetime.now(timezone.utc).isoformat()

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")


# ─────────────────────────── plumbing ───────────────────────────
def load_env():
    env = {}
    p = ROOT / ".env.local"
    if p.exists():
        for line in p.read_text(encoding="utf-8").splitlines():
            m = re.match(r"\s*([A-Z0-9_]+)\s*=\s*(.*)\s*$", line)
            if m:
                env[m.group(1)] = m.group(2).strip().strip('"').strip("'")
    env.update({k: v for k, v in os.environ.items() if k.startswith(("NEXT_PUBLIC_SUPABASE", "SUPABASE_"))})
    return env


class _Result:
    def __init__(self, data):
        self.data = data


class _Query:
    """The slice of the PostgREST query builder this script needs, over plain requests."""

    def __init__(self, client, table):
        self.c, self.table, self.params, self.method, self.body, self.prefer = client, table, [], "GET", None, []

    def select(self, cols="*"):
        self.params.append(("select", "".join(cols.split())))
        return self

    def eq(self, col, val):
        self.params.append((col, f"eq.{str(val).lower() if isinstance(val, bool) else val}"))
        return self

    def in_(self, col, vals):
        quoted = ",".join('"' + str(v).replace('"', '\\"') + '"' for v in vals)
        self.params.append((col, f"in.({quoted})"))
        return self

    def order(self, col, desc=False):
        self.params.append(("order", f"{col}.{'desc' if desc else 'asc'}"))
        return self

    def limit(self, n):
        self.params.append(("limit", str(n)))
        return self

    def range(self, start, end):
        self.params += [("offset", str(start)), ("limit", str(end - start + 1))]
        return self

    def update(self, values):
        self.method, self.body = "PATCH", values
        self.prefer.append("return=minimal")
        return self

    def insert(self, rows):
        self.method, self.body = "POST", rows
        self.prefer.append("return=representation")
        return self

    def upsert(self, row, on_conflict):
        self.method, self.body = "POST", row
        self.params.append(("on_conflict", on_conflict))
        self.prefer += ["resolution=merge-duplicates", "return=minimal"]
        return self

    def execute(self):
        import requests
        headers = dict(self.c.headers)
        if self.prefer:
            headers["Prefer"] = ",".join(self.prefer)
        r = requests.request(self.method, f"{self.c.url}/rest/v1/{self.table}", params=self.params,
                             json=self.body, headers=headers, timeout=60)
        if r.status_code >= 400:
            raise RuntimeError(f"{self.method} {self.table}: {r.status_code} {r.text[:300]}")
        return _Result(r.json() if r.content else None)


class _Client:
    def __init__(self, url, key):
        self.url = url.rstrip("/")
        self.headers = {"apikey": key, "Authorization": f"Bearer {key}", "Content-Type": "application/json"}

    def table(self, name):
        return _Query(self, name)


def db():
    env = load_env()
    url, key = env.get("NEXT_PUBLIC_SUPABASE_URL"), env.get("SUPABASE_SERVICE_ROLE_KEY")
    if not url or not key:
        sys.exit("NEXT_PUBLIC_SUPABASE_URL and SUPABASE_SERVICE_ROLE_KEY are required (.env.local).")
    return _Client(url, key)


def fetch_by_ids(sb, table, ids, cols="*"):
    out = {}
    ids = [i for i in dict.fromkeys(ids) if i and UUID_RE.match(i)]  # an edited id is held later, not queried
    for k in range(0, len(ids), 100):
        for r in sb.table(table).select(cols).in_("id", ids[k:k + 100]).execute().data or []:
            out[r["id"]] = r
    return out


def fetch_where_in(sb, table, col, values, cols="*"):
    out = []
    values = [v for v in dict.fromkeys(values) if v]
    for k in range(0, len(values), 100):
        start = 0
        while True:
            page = (sb.table(table).select(cols).in_(col, values[k:k + 100])
                    .range(start, start + 999).execute().data or [])
            out += page
            if len(page) < 1000:
                break
            start += 1000
    return out


def cell(v):
    if v is None:
        return ""
    return str(v).strip()


def norm(s):
    return " ".join(cell(s).split())


UUID_RE = re.compile(r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$", re.I)


def read_sheet(ws, marker=VERDICT):
    """Find the header row (it carries the verdict or the id column) and return rows keyed by header."""
    for r in range(1, min(ws.max_row, 8) + 1):
        heads = [cell(c.value) for c in ws[r]]
        if marker in heads or ID in heads:
            cols = {h: i for i, h in enumerate(heads) if h}
            rows = []
            for rn in range(r + 1, ws.max_row + 1):
                vals = [c.value for c in ws[rn]]
                rows.append((rn, {h: cell(vals[i]) if i < len(vals) else "" for h, i in cols.items()}))
            return list(cols), rows
    return None, []


class Plan:
    """Everything one workbook would do, before anything is written."""

    def __init__(self, path, bank, reviewer):
        self.path, self.bank, self.reviewer = Path(path), bank, reviewer
        self.sha = hashlib.sha256(Path(path).read_bytes()).hexdigest()
        self.changes = []   # dicts: table, id, op, set, before, sheet, row, verdict, action, review, label
        self.held = []      # dicts: sheet, row, id, verdict, reason, review, label, table
        self.notes = []     # dicts: sheet, row, id, review, label
        self.warnings = []  # (sheet, row, id, text)
        self.coverage = []  # plain sentences for the summary
        self.unreviewed = 0
        self.cut_score = None

    def hold(self, sheet, row, item_id, verdict, reason, review, label="", table=None):
        self.held.append(dict(sheet=sheet, row=row, id=item_id, verdict=verdict, reason=reason,
                              review=review, label=label, table=table))

    def change(self, table, item_id, op, values, before, sheet, row, verdict, action, review, label=""):
        self.changes.append(dict(table=table, id=item_id, op=op, set=values, before=before, sheet=sheet, row=row,
                                 verdict=verdict, action=action, review=review, label=label))

    def warn(self, sheet, row, item_id, text):
        self.warnings.append((sheet, row, item_id, text))


def review_cols(rowd, source):
    skip = set(source) | {"#", ID}
    return {h: v for h, v in rowd.items() if h not in skip and not h.startswith("_") and v}


def verdict_of(rowd, header=VERDICT):
    v = cell(rowd.get(header)).lower()
    return {"approve": "approve", "revise": "revise", "reject": "reject"}.get(v, v)


def is_no(v):
    return cell(v).lower() in ("no", "out of date", "none of them")


# ─────────────────────────── shared row logic ───────────────────────────
def iter_reviewed(plan, ws, source, en_header, table, bank_rows, strip_tip=False, id_required=True, rows=None):
    """Yield (rowno, rowd, item_id, db_row, verdict, review) for reviewed rows, holding the unsafe ones."""
    if rows is None:
        heads, rows = read_sheet(ws)
        if not heads:
            return
    seen = Counter(r[ID] for _, r in rows if r.get(ID))
    for rn, rowd in rows:
        rev = review_cols(rowd, source)
        item_id = rowd.get(ID, "")
        text = norm(rowd.get(en_header))
        verdict = verdict_of(rowd)
        label = text[:90]
        if not item_id and not id_required:  # an authoring row (new AC rating anchor)
            if text or cell(rowd.get("Suggested wording")):
                yield rn, rowd, None, None, verdict, rev
            continue
        if not item_id and not rev and not text:
            continue
        if not rev:
            plan.unreviewed += 1
            continue
        if not item_id:
            plan.hold(ws.title, rn, None, verdict, "No Item ID on this row, so it cannot be matched to the bank",
                      rev, label, table)
            continue
        if not UUID_RE.match(item_id):
            plan.hold(ws.title, rn, item_id, verdict, "Item ID has been edited and is no longer valid", rev, label, table)
            continue
        if seen[item_id] > 1:
            plan.hold(ws.title, rn, item_id, verdict, "The same Item ID appears on more than one row", rev, label, table)
            continue
        row = bank_rows.get(item_id)
        if row is None:
            plan.hold(ws.title, rn, item_id, verdict, f"Item ID not found in {table}", rev, label, table)
            continue
        if verdict not in ("approve", "revise", "reject"):
            if verdict:
                plan.hold(ws.title, rn, item_id, verdict, f"Unrecognised verdict '{rowd.get(VERDICT)}'", rev, label, table)
            else:
                plan.notes.append(dict(sheet=ws.title, row=rn, id=item_id, review=rev, label=label))
            continue
        if verdict == "approve" and any(is_no(v) for h, v in rev.items() if h.startswith(("Measures", "Develops"))):
            plan.warn(ws.title, rn, item_id, "Approved, although the reviewer says it does not measure its construct")
        yield rn, rowd, item_id, row, verdict, rev


def drifted(db_text, sheet_text, strip_tip=False):
    a = cell(db_text)
    if strip_tip:
        a = a.replace(TIP, "", 1)
    return norm(a) != norm(sheet_text)


def text_revision(plan, sheet, rn, item_id, rowd, rev, label, table, row, en_db, ar_db,
                  en_sugg="Suggested wording (EN)", ar_sugg="Suggested wording (AR)"):
    """Return {en_db: ..., ar_db: ...} for a Revise, or None after holding it."""
    s_en, s_ar = cell(rowd.get(en_sugg)), cell(rowd.get(ar_sugg)) if ar_sugg else ""
    if not s_en and not s_ar:
        plan.hold(sheet, rn, item_id, "revise", "Revise with no replacement wording", rev, label, table)
        return None
    has_ar = bool(ar_db and cell(row.get(ar_db)))
    if s_en and has_ar and not s_ar:
        plan.hold(sheet, rn, item_id, "revise",
                  "English revised without matching Arabic; the Arabic would no longer say the same thing",
                  rev, label, table)
        return None
    out = {}
    if s_en:
        out[en_db] = s_en
    if s_ar and ar_db:
        out[ar_db] = s_ar
    return out


# ─────────────────────────── banks ───────────────────────────
def plan_persona(sb, wb, plan):
    ws = wb["Items"]
    _, rows = read_sheet(ws)
    bank = fetch_by_ids(sb, "persona_items", [r.get(ID) for _, r in rows])
    src = ["Competency", "Competency (AR)", "Item", "Reverse", "Statement (EN)", "Statement (AR)", "Bank status"]
    rejects = []
    comp_name = {bank[r[ID]]["ac_competency_id"]: r.get("Competency") for _, r in rows if r.get(ID) in bank}
    for rn, rowd, iid, row, verdict, rev in iter_reviewed(plan, ws, src, "Statement (EN)", "persona_items", bank):
        label = norm(rowd.get("Statement (EN)"))[:90]
        if drifted(row["text_en"], rowd.get("Statement (EN)")):
            plan.hold(ws.title, rn, iid, verdict, "Bank text changed after this workbook was issued", rev, label, "persona_items")
            continue
        stamp = {"sme_reviewer_name": plan.reviewer, "sme_reviewed_at": NOW}
        if verdict == "approve":
            if is_no(rowd.get("Reverse flag correct?")):
                plan.hold(ws.title, rn, iid, verdict, "Reviewer says the reverse-keying flag is wrong; changing it "
                          "changes how past sittings were scored", rev, label, "persona_items")
                continue
            vals = {"status": "approved", **stamp}
            if cell(rowd.get("Arabic OK?")).lower() == "yes":
                vals["ar_reviewed"] = True
            elif is_no(rowd.get("Arabic OK?")):
                plan.warn(ws.title, rn, iid, "Approved, but the Arabic was marked not OK with no replacement given")
            plan.change("persona_items", iid, "update", vals, _pick(row, vals), ws.title, rn, verdict, "approve", rev, label)
        elif verdict == "revise":
            txt = text_revision(plan, ws.title, rn, iid, rowd, rev, label, "persona_items", row, "text_en", "text_ar")
            if txt is None:
                continue
            vals = {**txt, "status": "approved", **stamp}
            if "text_ar" in txt:
                vals["ar_reviewed"] = True
            plan.change("persona_items", iid, "update", vals, _pick(row, vals), ws.title, rn, verdict, "revise", rev, label)
        else:
            rejects.append((rn, iid, row, rev, label))
            plan.change("persona_items", iid, "update", {"status": "rejected", **stamp},
                        _pick(row, {"status": 1}), ws.title, rn, verdict, "reject", rev, label)

    # A rejection must not leave a competency unscoreable.
    if rejects:
        comps = {r["ac_competency_id"] for _, _, r, _, _ in rejects}
        allrows = fetch_where_in(sb, "persona_items", "ac_competency_id", comps,
                                 "id, ac_competency_id, reverse, status")
        rejected_ids = {iid for _, iid, _, _, _ in rejects}
        for comp in comps:
            live = [r for r in allrows if r["ac_competency_id"] == comp and r["status"] in ("pending", "approved")
                    and r["id"] not in rejected_ids]
            ok = len(live) >= 3 and any(r["reverse"] for r in live)
            if not ok:
                for rn, iid, r, rev, label in rejects:
                    if r["ac_competency_id"] == comp:
                        _unchange(plan, iid)
                        plan.hold(ws.title, rn, iid, "reject",
                                  f"Rejecting would leave this competency with {len(live)} live statements"
                                  + ("" if any(x['reverse'] for x in live) else " and no reverse-keyed statement")
                                  + "; a replacement must be written first", rev, label, "persona_items")
            else:
                plan.coverage.append(f"Persona {comp_name.get(comp, comp)}: {len(live)} live statements after rejections")
    notes_sheet(plan, wb, "Competency review", "Competency")


def plan_reflect(sb, wb, plan):
    ws = wb["Behaviours"]
    _, rows = read_sheet(ws)
    bank = fetch_by_ids(sb, "reflect_behaviors", [r.get(ID) for _, r in rows])
    src = ["Competency", "Competency (AR)", "What the competency means", "Statement (EN)", "Statement (AR)"]
    for rn, rowd, iid, row, verdict, rev in iter_reviewed(plan, ws, src, "Statement (EN)", "reflect_behaviors", bank):
        label = norm(rowd.get("Statement (EN)"))[:90]
        if drifted(row["text_en"], rowd.get("Statement (EN)")):
            plan.hold(ws.title, rn, iid, verdict, "Bank text changed after this workbook was issued", rev, label,
                      "reflect_behaviors")
            continue
        stamp = {"sme_reviewer_name": plan.reviewer, "sme_reviewed_at": NOW}
        if verdict == "approve":
            vals = {"sme_status": "approved", **stamp}
            if is_no(rowd.get("Arabic OK?")):
                plan.warn(ws.title, rn, iid, "Approved, but the Arabic was marked not OK with no replacement given")
        elif verdict == "revise":
            txt = text_revision(plan, ws.title, rn, iid, rowd, rev, label, "reflect_behaviors", row, "text_en", "text_ar")
            if txt is None:
                continue
            vals = {**txt, "sme_status": "approved", **stamp}
        else:
            vals = {"sme_status": "rejected", **stamp}
            plan.warn(ws.title, rn, iid, "Rejected statements stay in the library until removed in Reflect admin; "
                                         "engagements already running keep their own copies")
        plan.change("reflect_behaviors", iid, "update", vals, _pick(row, vals), ws.title, rn, verdict, verdict, rev, label)
    notes_sheet(plan, wb, "Competency review", "Competency")


def plan_ac(sb, wb, plan):
    comps = {c["name"]: c["id"] for c in sb.table("competencies").select("id, name").execute().data or []}
    next_sort = {}
    for sheet, en_header, is_tip in (("Rating anchors", "Indicator (EN)", False), ("Development tips", "Development tip", True)):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        _, rows = read_sheet(ws)
        bank = fetch_by_ids(sb, "behavioral_indicators", [r.get(ID) for _, r in rows])
        src = ["Competency", "Type", en_header]
        last_comp = ""
        for rn, rowd in rows:  # carry the competency down for the authoring rows
            if rowd.get("Competency"):
                last_comp = rowd["Competency"]
            rowd["_comp"] = last_comp
        for rn, rowd, iid, row, verdict, rev in iter_reviewed(plan, ws, src, en_header, "behavioral_indicators", bank,
                                                              strip_tip=is_tip, id_required=is_tip, rows=rows):
            label = norm(rowd.get(en_header) or rowd.get("Suggested wording"))[:90]
            stamp = {"sme_reviewer_name": plan.reviewer, "sme_reviewed_at": NOW}
            if iid is None:  # a new rating anchor written by the reviewer
                text = cell(rowd.get(en_header)) or cell(rowd.get("Suggested wording"))
                typ = cell(rowd.get("Type")).lower()
                comp_id = comps.get(rowd["_comp"])
                if verdict == "reject" or not text:
                    continue
                if not comp_id:
                    plan.hold(sheet, rn, None, verdict, f"New anchor: competency '{rowd['_comp']}' not recognised", rev,
                              label, "behavioral_indicators")
                    continue
                if typ not in ("positive", "contra-indicator"):
                    plan.hold(sheet, rn, None, verdict, "New anchor: choose Positive or Contra-indicator in the Type column",
                              rev, label, "behavioral_indicators")
                    continue
                if comp_id not in next_sort:
                    existing = sb.table("behavioral_indicators").select("sort_order").eq("competency_id", comp_id) \
                        .order("sort_order", desc=True).limit(1).execute().data
                    next_sort[comp_id] = (existing[0]["sort_order"] if existing else 0) + 1
                vals = {"competency_id": comp_id, "indicator_type": "negative" if typ == "contra-indicator" else "positive",
                        "description": text, "sort_order": next_sort[comp_id], "sme_status": "approved", **stamp}
                next_sort[comp_id] += 1
                plan.change("behavioral_indicators", None, "insert", vals, None, sheet, rn, verdict or "new", "insert", rev, label)
                continue
            if drifted(row["description"], rowd.get(en_header), strip_tip=is_tip):
                plan.hold(sheet, rn, iid, verdict, "Bank text changed after this workbook was issued", rev, label,
                          "behavioral_indicators")
                continue
            vals = {}
            if not is_tip and is_no(rowd.get("Type correct?")):
                flipped = "negative" if row["indicator_type"] == "positive" else "positive"
                vals["indicator_type"] = flipped
                plan.warn(sheet, rn, iid, f"Type corrected to {'Contra-indicator' if flipped == 'negative' else 'Positive'}")
            if verdict == "approve":
                vals.update({"sme_status": "approved", **stamp})
            elif verdict == "revise":
                s = cell(rowd.get("Suggested wording"))
                if not s:
                    plan.hold(sheet, rn, iid, verdict, "Revise with no replacement wording", rev, label, "behavioral_indicators")
                    continue
                vals.update({"description": f"{TIP} {s}" if is_tip else s, "sme_status": "approved", **stamp})
            else:
                vals.update({"sme_status": "rejected", **stamp})
                plan.warn(sheet, rn, iid, "Rejected indicators stay visible to assessors until removed from the competency")
            plan.change("behavioral_indicators", iid, "update", vals, _pick(row, vals), sheet, rn, verdict, verdict, rev, label)


def plan_arc(sb, wb, plan):
    if "Maturity ladders" in wb.sheetnames:  # cross-pillar scoring pack: rulings only, nothing to write
        for sheet in wb.sheetnames:
            if sheet == "Instructions":
                continue
            heads, rows = read_sheet(wb[sheet], marker="Comments")
            for rn, rowd in rows:
                rev = {h: v for h, v in rowd.items() if v}
                ruling_cols = [h for h in (heads or []) if h not in ("#", "Pillar", "Q no.", "Question", "Option 1",
                               "Option 2", "Option 3", "Option 4", "Option 5", "Scores (1-5)", "Scores", ID, "Finding",
                               "Where", "Scored items", "Likert", "Multiple choice", "Yes/No")]
                if any(rowd.get(h) for h in ruling_cols):
                    plan.notes.append(dict(sheet=sheet, row=rn, id=rowd.get(ID) or None, review=rev,
                                           label=norm(rowd.get("Question") or rowd.get("Finding") or rowd.get("Pillar"))[:90]))
        return
    ids = []
    for sheet in ("Items", "Scenario items", "Consultant probes"):
        if sheet in wb.sheetnames:
            ids += [r.get(ID) for _, r in read_sheet(wb[sheet])[1]]
    bank = fetch_by_ids(sb, "ara_questions", ids)
    for sheet, en_header, src in (
        ("Items", "Question (EN)", ["Q no.", "Pillar", "Factor", "Dimension", "Layer", "Type", "Question (EN)",
                                    "Question (AR)", "Response options", "How it scores", "Anchor / construct"]),
        ("Scenario items", "Scenario (EN)", ["Q no.", "Factor", "Type", "Scenario (EN)", "Scenario (AR)", "Options (EN)",
                                             "Keyed best response", "Score"]),
        ("Consultant probes", "Question (EN)", ["Q no.", "Pillar", "Question (EN)", "Question (AR)", "Reference note"]),
    ):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        for rn, rowd, iid, row, verdict, rev in iter_reviewed(plan, ws, src, en_header, "ara_questions", bank):
            label = norm(rowd.get(en_header))[:90]
            if drifted(row["question_text_en"], rowd.get(en_header)):
                plan.hold(sheet, rn, iid, verdict, "Bank text changed after this workbook was issued", rev, label, "ara_questions")
                continue
            stamp = {"sme_reviewer_name": plan.reviewer, "sme_reviewed_at": NOW}
            if verdict == "approve":
                if sheet == "Scenario items" and is_no(rowd.get("Key is right?")):
                    plan.hold(sheet, rn, iid, verdict, "Approved, but the reviewer says the keyed response is wrong", rev,
                              label, "ara_questions")
                    continue
                vals = {"sme_status": "approved", **stamp}
            elif verdict == "revise":
                # v1.1 is never reworded in place; the v1.2 build applies these from the log.
                if sheet != "Items":
                    detail = "free-text fix to a scenario or probe, to be written into v1.2 by hand"
                elif not (rowd.get("Suggested wording (EN)") or rowd.get("Suggested wording (AR)")):
                    detail = "no replacement wording given; the reviewer's comments guide the v1.2 edit"
                elif rowd.get("Suggested wording (EN)") and row.get("question_text_ar") and not rowd.get("Suggested wording (AR)"):
                    detail = "English replacement given; Arabic still to be translated for v1.2"
                else:
                    detail = "replacement wording given"
                plan.hold(sheet, rn, iid, verdict, f"{ARC_V12_QUEUE}: {detail}", rev, label, "ara_questions")
                continue
            else:
                vals = {"sme_status": "rejected", **stamp}
                plan.warn(sheet, rn, iid, "Still served on v1.1 (results stay provisional); left out of v1.2")
            plan.change("ara_questions", iid, "update", vals, _pick(row, vals), sheet, rn, verdict, verdict, rev, label)
    for s in ("Pillar review", "Factor review", "Dimension review"):
        notes_sheet(plan, wb, s, s.split()[0])


def _key_checks(plan, sheet, rn, iid, rowd, rev, label, table, key_header="Key"):
    """Hold an Approve that contradicts the reviewer's own answer checks. Returns True if held."""
    if is_no(rowd.get("Key correct?")) or is_no(rowd.get("Key is right?")):
        plan.hold(sheet, rn, iid, "approve", "Approved, but the reviewer says the key is wrong", rev, label, table)
        return True
    mine, key = cell(rowd.get("Your answer")).upper(), cell(rowd.get(key_header)).upper()
    if mine == "NONE OF THEM":
        plan.hold(sheet, rn, iid, "approve", "Approved, but the reviewer found none of the options correct", rev, label, table)
        return True
    if mine and len(key) == 1 and mine != key:
        plan.hold(sheet, rn, iid, "approve", f"Approved, but the reviewer answered {mine} and the key is {key}", rev,
                  label, table)
        return True
    if cell(rowd.get("Current practice?")).lower() == "out of date":
        plan.hold(sheet, rn, iid, "approve", "Approved, but marked out of date under current practice", rev, label, table)
        return True
    return False


def plan_technical(sb, wb, plan):
    ws = wb["Items"]
    _, rows = read_sheet(ws)
    bank = fetch_by_ids(sb, "tech_assessment_items", [r.get(ID) for _, r in rows])
    src = ["Skill", "Difficulty", "Type", "Scenario (EN)", "Question (EN)", "A", "B", "C", "D", "Key",
           "Author explanation", "Question (AR)", "A (AR)", "B (AR)", "C (AR)", "D (AR)", "Bank status"]
    for rn, rowd, iid, row, verdict, rev in iter_reviewed(plan, ws, src, "Question (EN)", "tech_assessment_items", bank):
        label = norm(rowd.get("Question (EN)"))[:90]
        if drifted(row["question_en"], rowd.get("Question (EN)")):
            plan.hold(ws.title, rn, iid, verdict, "Bank text changed after this workbook was issued", rev, label,
                      "tech_assessment_items")
            continue
        note = f"SME review ({plan.reviewer}, {NOW[:10]}): {verdict}" + (
            f" - {rowd['Comments']}" if rowd.get("Comments") else "")
        notes = (cell(row.get("review_notes")) + "\n" + note).strip()
        stamp = {"reviewer_name": plan.reviewer, "reviewed_at": NOW, "review_notes": notes}
        if verdict == "approve":
            if _key_checks(plan, ws.title, rn, iid, rowd, rev, label, "tech_assessment_items"):
                continue
            vals = {"status": "approved", **stamp}
        elif verdict == "revise":
            plan.hold(ws.title, rn, iid, verdict, "Technical fixes (stem, options or key) need a person to make them in "
                      "the item review console: " + (rowd.get("Suggested fix") or "no fix written"), rev, label,
                      "tech_assessment_items")
            continue
        else:
            vals = {"status": "rejected", **stamp}
        plan.change("tech_assessment_items", iid, "update", vals, _pick(row, vals), ws.title, rn, verdict, verdict, rev, label)

    # Cut-score sheet
    domains = Counter(r["domain_key"] for r in bank.values())
    if "Cut-score" in wb.sheetnames and domains:
        domain = domains.most_common(1)[0][0]
        fields = {}
        for r in wb["Cut-score"].iter_rows(values_only=True):
            if r and cell(r[0]):
                fields[cell(r[0])] = cell(r[1]) if len(r) > 1 else ""
        pass_raw = fields.get("Recommended pass mark (%)", "")
        if pass_raw:
            try:
                pass_pct = float(pass_raw.rstrip("%"))
                min_items = int(float(fields.get("Minimum items a test must draw") or 8))
                if not (1 <= pass_pct <= 100) or min_items < 8:
                    raise ValueError
                rationale = "; ".join(x for x in (
                    fields.get("Your reasoning"),
                    f"Mandatory skills: {fields['Should any skill be mandatory?']}"
                    if fields.get("Should any skill be mandatory?") else "") if x)
                plan.cut_score = {"domain_key": domain, "pass_pct": pass_pct, "min_items": min_items,
                                  "method": "SME recommendation (VIFM review workbook)",
                                  "rationale": rationale or None,
                                  "set_by_name": fields.get("Recommended by") or plan.reviewer,
                                  "set_at": NOW, "updated_at": NOW}
            except ValueError:
                plan.hold("Cut-score", None, None, "", f"Pass mark '{pass_raw}' or minimum items is not usable "
                          "(pass mark 1-100, minimum items 8 or more)", fields, domain, "tech_assessment_cut_scores")

    # Certification readiness after this import
    for domain in domains:
        allrows = sb.table("tech_assessment_items").select("id, status, skill").eq("domain_key", domain).execute().data or []
        new_status = {c["id"]: c["set"]["status"] for c in plan.changes if c["table"] == "tech_assessment_items"}
        approved = [r for r in allrows if new_status.get(r["id"], r["status"]) == "approved"]
        served = [r for r in allrows if new_status.get(r["id"], r["status"]) in ("approved", "in_review")]
        empty = sorted({r["skill"] for r in allrows} - {r["skill"] for r in served})
        plan.coverage.append(f"Technical {domain}: {len(approved)} approved after import "
                             f"(certified tests need {plan.cut_score['min_items'] if plan.cut_score else 8})"
                             + (f"; skills left with no servable item: {', '.join(empty)}" if empty else ""))


def plan_logica(sb, wb, plan, drafts=False):
    ws = wb["Items"]
    _, rows = read_sheet(ws)
    bank = fetch_by_ids(sb, "psy_items", [r.get(ID) for _, r in rows])
    facet_map = {"ratio and proportion": "num_ratio", "percentages": "num_percent", "data interpretation": "num_data"}
    src = (["Stem (EN)", "Stem (AR)", "Option A", "Option B", "Option C", "Option D", "Correct", "Drafted difficulty"]
           if drafts else
           ["Facet", "Difficulty", "Stem (EN)", "A", "B", "C", "D", "Key", "Author rationale", "Stem (AR)", "A (AR)",
            "B (AR)", "C (AR)", "D (AR)", "Bank status"])
    rejects = []
    for rn, rowd, iid, row, verdict, rev in iter_reviewed(plan, ws, src, "Stem (EN)", "psy_items", bank):
        label = norm(rowd.get("Stem (EN)"))[:90]
        if drifted(row["stem_en"], rowd.get("Stem (EN)")):
            plan.hold(ws.title, rn, iid, verdict, "Bank text changed after this workbook was issued", rev, label, "psy_items")
            continue
        stamp = {"sme_reviewer_name": plan.reviewer, "reviewed_at": NOW}
        if verdict == "approve":
            if _key_checks(plan, ws.title, rn, iid, rowd, rev, label, "psy_items",
                           key_header="Correct" if drafts else "Key"):
                continue
            vals = {"status": "approved", **stamp}
            if drafts:
                facet = facet_map.get(cell(rowd.get("Facet")).lower())
                diff = cell(rowd.get("Difficulty")).lower()
                if not facet or diff not in ("easy", "medium", "hard"):
                    plan.hold(ws.title, rn, iid, verdict, "Approved draft needs both a Facet and a Difficulty before it can "
                              "be served", rev, label, "psy_items")
                    continue
                vals.update({"facet": facet, "difficulty": diff})
            if cell(rowd.get("Arabic OK?")).lower() == "yes":
                vals["ar_reviewed"] = True
            if cell(rowd.get("Rule stated in stem?")).lower() == "yes":
                plan.warn(ws.title, rn, iid, "Approved although the reviewer says the stem states the rule")
            if is_no(rowd.get("Difficulty agree?")):
                plan.warn(ws.title, rn, iid, "Reviewer disagrees with the difficulty label; see Comments")
        elif verdict == "revise":
            plan.hold(ws.title, rn, iid, verdict, "Logica fixes (stem, options or key) need a person to make them: "
                      + (rowd.get("Suggested fix") or rowd.get("Suggested wording") or "no fix written"), rev, label,
                      "psy_items")
            continue
        else:
            vals = {"status": "rejected", "rejected_reason": rowd.get("Comments") or "Rejected in SME review", **stamp}
            rejects.append((rn, iid, row, rev, label))
        plan.change("psy_items", iid, "update", vals, _pick(row, vals), ws.title, rn, verdict, verdict, rev, label)

    # A rejection must not empty a facet x difficulty cell: the test would stop assembling.
    if rejects:
        scales = {r["scale_id"] for _, _, r, _, _ in rejects}
        allrows = fetch_where_in(sb, "psy_items", "scale_id", scales, "id, scale_id, facet, difficulty, status")
        gone = {iid for _, iid, _, _, _ in rejects}
        cells = Counter((r["facet"], r["difficulty"]) for r in allrows
                        if r["status"] in ("approved", "in_review") and r["id"] not in gone and r["facet"])
        for rn, iid, r, rev, label in rejects:
            key = (r["facet"], r["difficulty"])
            if r["facet"] and cells[key] < 1:
                _unchange(plan, iid)
                plan.hold(ws.title, rn, iid, "reject", f"Rejecting would leave {r['facet']} / {r['difficulty']} with no "
                          "live question, and the whole subtest would stop assembling; a replacement must be approved "
                          "first", rev, label, "psy_items")
        for (facet, diff), n in sorted(cells.items()):
            if n <= 1:
                plan.coverage.append(f"Logica {facet} / {diff}: only {n} live question left after import")


def notes_sheet(plan, wb, sheet, first_col):
    if sheet not in wb.sheetnames:
        return
    _, rows = read_sheet(wb[sheet], marker=first_col)
    for rn, rowd in rows:
        rev = {h: v for h, v in rowd.items() if v and h not in (first_col, "Items")}
        if rev and rowd.get(first_col):
            plan.notes.append(dict(sheet=sheet, row=rn, id=None, review={first_col: rowd[first_col], **rev},
                                   label=rowd[first_col]))


def _pick(row, vals):
    return {k: row.get(k) for k in vals if k in row}


def _unchange(plan, iid):
    plan.changes = [c for c in plan.changes if c["id"] != iid]


# ─────────────────────────── dispatch ───────────────────────────
def detect(path, wb):
    name = Path(path).name
    if name.startswith("Technical-"):
        return "technical"
    if name.startswith("Logica-Numerical-Drafts"):
        return "logica_drafts"
    if name.startswith("Logica-"):
        return "logica"
    if name.startswith("Persona-"):
        return "persona"
    if name.startswith("ARC-"):
        return "arc"
    if name.startswith("AC-"):
        return "ac"
    if name.startswith("Reflect-360-"):
        return "reflect"
    if name.startswith("Role-Readiness"):
        return "role_readiness"
    if name.startswith("Fluent-"):
        return "fluent"
    return None


def reviewer_from_register(path):
    if not REGISTER.exists():
        return None
    # Tolerate mail-download copies ("file (1).xlsx") and a reviewer's suffix ("file-Ali.xlsx").
    stem = re.sub(r"\s*\(\d+\)$", "", Path(path).stem)
    ws = load_workbook(REGISTER, read_only=True)["Assignment plan"]
    best = None
    for row in ws.iter_rows(min_row=6, values_only=True):
        if row and len(row) > 8 and row[8]:
            reg = Path(str(row[8])).stem
            if stem == reg or stem.startswith(reg):
                if best is None or len(reg) > len(best[0]):
                    best = (reg, row[4])
    return best[1] if best else None


def build_plan(sb, path, reviewer):
    wb = load_workbook(path, data_only=True)
    bank = detect(path, wb)
    plan = Plan(path, bank or "unknown", reviewer)
    if bank == "persona":
        plan_persona(sb, wb, plan)
    elif bank == "reflect":
        plan_reflect(sb, wb, plan)
    elif bank == "ac":
        plan_ac(sb, wb, plan)
    elif bank == "arc":
        plan_arc(sb, wb, plan)
    elif bank == "technical":
        plan_technical(sb, wb, plan)
    elif bank == "logica":
        plan_logica(sb, wb, plan)
    elif bank == "logica_drafts":
        plan.bank = "logica"
        plan_logica(sb, wb, plan, drafts=True)
    else:
        plan.hold("-", None, None, "", f"No importer for this workbook type yet ({bank or 'unrecognised file name'})", {})
    return plan


# ─────────────────────────── output ───────────────────────────
def print_plan(plan):
    acts = Counter(c["action"] for c in plan.changes)
    print(f"\n{plan.path.name}")
    print(f"  bank: {plan.bank}   reviewer: {plan.reviewer}")
    print(f"  to apply: {len(plan.changes)}  (approve {acts['approve']}, revise {acts['revise']}, "
          f"reject {acts['reject']}, new {acts['insert']})" + ("  + cut-score" if plan.cut_score else ""))
    print(f"  held for a person: {len(plan.held)}   notes: {len(plan.notes)}   not yet reviewed: {plan.unreviewed}")
    for h in plan.held[:12]:
        print(f"    HELD {h['sheet']} row {h['row']}: {h['reason']}")
    if len(plan.held) > 12:
        print(f"    ... {len(plan.held) - 12} more in the report")
    for w in plan.warnings[:8]:
        print(f"    note {w[0]} row {w[1]}: {w[3]}")
    for c in plan.coverage:
        print(f"    coverage: {c}")
    if plan.cut_score:
        cs = plan.cut_score
        print(f"    cut-score: {cs['domain_key']} pass {cs['pass_pct']}% min items {cs['min_items']}")


def write_report(plans, out, applied):
    wb = Workbook()
    hdr_font, hdr_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="010131")
    wrap = Alignment(wrap_text=True, vertical="top")

    def sheet(title, headers, rows, widths):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill, c.alignment = hdr_font, hdr_fill, wrap
        for r in rows:
            ws.append([("" if v is None else v) for v in r])
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = wrap
        ws.freeze_panes = "A2"
        return ws

    ws = wb.active
    ws.title = "Summary"
    ws.append(["SME import " + ("(APPLIED)" if applied else "(PREVIEW - nothing written)"), NOW[:19].replace("T", " ")])
    ws["A1"].font = Font(bold=True, size=13, color="010131")
    ws.append([])
    ws.append(["Workbook", "Bank", "Reviewer", "To apply", "Approve", "Revise", "Reject", "New", "Cut-score", "Held",
               "Notes", "Not yet reviewed"])
    for c in ws[3]:
        c.font, c.fill = hdr_font, hdr_fill
    for p in plans:
        a = Counter(c["action"] for c in p.changes)
        ws.append([p.path.name, p.bank, p.reviewer, len(p.changes), a["approve"], a["revise"], a["reject"], a["insert"],
                   "yes" if p.cut_score else "", len(p.held), len(p.notes), p.unreviewed])
    ws.append([])
    for p in plans:
        for line in p.coverage:
            ws.append([p.path.name, line])
    ws.column_dimensions["A"].width = 58
    for col in "BCDEFGHIJKL":
        ws.column_dimensions[col].width = 13

    def fmt(d):
        return "\n".join(f"{k}: {v}" for k, v in (d or {}).items())

    sheet("To apply", ["Workbook", "Sheet", "Row", "Item ID", "Action", "Item", "Current", "New", "Reviewer's answers"],
          [[p.path.name, c["sheet"], c["row"], c["id"] or "(new)", c["action"], c["label"], fmt(c["before"]),
            fmt({k: v for k, v in c["set"].items() if not k.startswith(("sme_review", "review", "reviewed"))}),
            fmt(c["review"])] for p in plans for c in p.changes],
          [30, 16, 6, 38, 9, 50, 40, 40, 50])
    sheet("Held", ["Workbook", "Sheet", "Row", "Item ID", "Verdict", "Why it was held", "Item", "Reviewer's answers"],
          [[p.path.name, h["sheet"], h["row"], h["id"], h["verdict"], h["reason"], h["label"], fmt(h["review"])]
           for p in plans for h in p.held], [30, 16, 6, 38, 9, 60, 50, 60])
    sheet("Notes and warnings", ["Workbook", "Sheet", "Row", "Item ID", "Item", "Note"],
          [[p.path.name, w[0], w[1], w[2], "", w[3]] for p in plans for w in p.warnings]
          + [[p.path.name, n["sheet"], n["row"], n["id"], n["label"], fmt(n["review"])] for p in plans for n in p.notes],
          [30, 18, 6, 38, 40, 70])
    wb.save(out)


# ─────────────────────────── apply ───────────────────────────
def apply_plan(sb, plan, imported_by):
    base = dict(imported_by=imported_by, reviewer_name=plan.reviewer, bank=plan.bank,
                workbook_file=plan.path.name, workbook_sha256=plan.sha)
    logs, errors = [], 0
    for c in plan.changes:
        try:
            if c["op"] == "insert":
                res = sb.table(c["table"]).insert(c["set"]).execute().data
                c["id"] = res[0]["id"] if res else None
            else:
                sb.table(c["table"]).update(c["set"]).eq("id", c["id"]).execute()
            logs.append(dict(base, target_table=c["table"], item_id=c["id"], action=c["action"], verdict=c["verdict"],
                             before=c["before"], after=c["set"], review=c["review"], applied=True,
                             sheet=c["sheet"], row_number=c["row"]))
        except Exception as e:  # keep going; the failure is logged as held
            errors += 1
            logs.append(dict(base, target_table=c["table"], item_id=c["id"], action="held", verdict=c["verdict"],
                             before=c["before"], after=c["set"], review=c["review"], applied=False,
                             held_reason=f"Write failed: {e}", sheet=c["sheet"], row_number=c["row"]))
    if plan.cut_score:
        try:
            sb.table("tech_assessment_cut_scores").upsert(plan.cut_score, on_conflict="domain_key").execute()
            logs.append(dict(base, target_table="tech_assessment_cut_scores", action="cut_score", after=plan.cut_score,
                             applied=True, sheet="Cut-score"))
        except Exception as e:
            errors += 1
            logs.append(dict(base, target_table="tech_assessment_cut_scores", action="held", after=plan.cut_score,
                             applied=False, held_reason=f"Write failed: {e}", sheet="Cut-score"))
    for h in plan.held:
        logs.append(dict(base, target_table=h.get("table"), item_id=h["id"], action="held", verdict=h["verdict"],
                         review=h["review"], held_reason=h["reason"], applied=False, sheet=h["sheet"],
                         row_number=h["row"]))
    for n in plan.notes:
        logs.append(dict(base, item_id=n["id"], action="note", review=n["review"], applied=False, sheet=n["sheet"],
                         row_number=n["row"]))
    for k in range(0, len(logs), 200):
        sb.table("sme_review_log").insert(logs[k:k + 200]).execute()
    return errors


def main():
    ap = argparse.ArgumentParser(description="Load returned SME review workbooks back into the banks.")
    ap.add_argument("paths", nargs="+", help="returned .xlsx workbooks, or folders of them")
    ap.add_argument("--reviewer", help="reviewer's name (defaults to the register's 'Assigned to')")
    ap.add_argument("--apply", action="store_true", help="write the changes (default is preview only)")
    ap.add_argument("--report", help="where to save the Excel report")
    ap.add_argument("--reimport", action="store_true", help="allow a workbook that was already applied to be applied again")
    args = ap.parse_args()

    files = []
    for p in args.paths:
        p = Path(p)
        files += sorted(p.glob("*.xlsx")) if p.is_dir() else [p]
    files = [f for f in files if not f.name.startswith("~$")]
    if not files:
        sys.exit("No .xlsx workbooks found.")

    sb = db()
    plans = []
    for f in files:
        reviewer = args.reviewer or reviewer_from_register(f)
        if not reviewer or reviewer in ("TO ASSIGN", "On hold", "No central review"):
            sys.exit(f"{f.name}: no reviewer found in the register; pass --reviewer \"Name\".")
        plan = build_plan(sb, f, reviewer)
        plans.append(plan)
        print_plan(plan)

    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    out = Path(args.report) if args.report else ROOT / ".tmp" / "sme" / "imports" / (
        f"sme-import-{'applied' if args.apply else 'preview'}-{stamp}.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)

    if args.apply:
        try:
            sb.table("sme_review_log").select("id").limit(1).execute()
        except Exception:
            sys.exit("sme_review_log is missing: apply migration 00203 before using --apply.")
        imported_by = f"sme_import.py ({getpass.getuser()})"
        total_errors = 0
        for plan in plans:
            done = sb.table("sme_review_log").select("id").eq("workbook_sha256", plan.sha).eq("applied", True) \
                .limit(1).execute().data
            if done and not args.reimport:
                print(f"\n{plan.path.name}: this exact workbook was already applied; skipped (use --reimport to force).")
                continue
            total_errors += apply_plan(sb, plan, imported_by)
        write_report(plans, out, applied=True)
        print(f"\nApplied. Report: {out}" + (f"   ({total_errors} write errors, logged as held)" if total_errors else ""))
    else:
        write_report(plans, out, applied=False)
        print(f"\nPreview only - nothing was written. Report: {out}")
        print("Rerun with --apply to write the 'To apply' changes.")


if __name__ == "__main__":
    main()
